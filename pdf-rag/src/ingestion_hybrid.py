#!/usr/bin/env python3
"""Гибридная индексация с поддержкой инкрементального добавления и метаданных продукта."""

import argparse
import hashlib
import json
import logging
import os
import re
import shutil
import sys
import pickle
from datetime import datetime, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

DATA_RAW = Path("data/pdfs_raw")
DATA_PDFS = Path("data/pdfs")
CHROMA_DIR = Path("data/chroma")
CHUNK_SIZE = 768
CHUNK_OVERLAP = 128
COLLECTION_NAME = "docs"
HASH_CACHE = CHROMA_DIR / "file_hashes.json"

# — Логирование
LOG_DIR = Path(__file__).resolve().parent.parent / "data" / "logs"
LOG_DIR.mkdir(parents=True, exist_ok=True)
INGEST_LOG = str(LOG_DIR / "ingestion.jsonl")

_logger = logging.getLogger("pdf-rag-ingest")
_logger.setLevel(logging.INFO)
_fh = logging.FileHandler(str(LOG_DIR / "ingestion.log"))
_fh.setFormatter(logging.Formatter("%(asctime)s [%(levelname)s] %(message)s"))
_logger.addHandler(_fh)


def log_ingestion(event: str, **kwargs):
    entry = {"ts": datetime.now(timezone.utc).isoformat(), "event": event}
    entry.update(kwargs)
    with open(INGEST_LOG, "a") as f:
        f.write(json.dumps(entry, ensure_ascii=False) + "\n")
    _logger.info(f"{event} {kwargs}")


def file_hash(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        for chunk in iter(lambda: f.read(65536), b''):
            h.update(chunk)
    return h.hexdigest()


def deduplicate_to_flat(target: Path = DATA_PDFS):
    """Копирует все файлы в плоскую структуру, удаляя дубликаты по SHA256."""
    target.mkdir(parents=True, exist_ok=True)
    seen = set()
    copied = skipped = 0
    for f in sorted(DATA_RAW.rglob('*')):
        if not f.is_file():
            continue
        ext = f.suffix.lower()
        if ext not in ('.pdf', '.docx', '.doc', '.pptx', '.zip'):
            continue
        h = file_hash(f)
        if h in seen:
            skipped += 1
            continue
        seen.add(h)
        stem = re.sub(r'[^\x20-\x7E\s\w\-.а-яёА-ЯЁ]', '_', f.stem)
        out = target / f"{stem}{f.suffix}"
        if out.exists():
            out = target / f"{stem}_{copied}{f.suffix}"
        shutil.copy2(f, out)
        copied += 1
    print(f"Дедупликация: {copied} скопировано, {skipped} пропущено (дубликаты)")
    return copied


def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == '.pdf':
        try:
            from pypdf import PdfReader
            r = PdfReader(str(path))
            return '\n'.join(p.extract_text() or '' for p in r.pages)
        except Exception as e:
            print(f"  [WARN] PDF error {path.name}: {e}")
            return ''
    elif ext == '.docx':
        try:
            from docx import Document
            d = Document(str(path))
            return '\n'.join(p.text for p in d.paragraphs)
        except Exception as e:
            print(f"  [WARN] DOCX error {path.name}: {e}")
            return ''
    elif ext == '.doc':
        print(f"  [SKIP] .doc (legacy format): {path.name}")
        return ''
    elif ext == '.pptx':
        try:
            from pptx import Presentation
            prs = Presentation(str(path))
            texts = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        texts.append(shape.text)
            return '\n'.join(texts)
        except Exception as e:
            print(f"  [WARN] PPTX error {path.name}: {e}")
            return ''
    return ''


def split_text(text: str, source_name: str) -> list[tuple[str, int]]:
    """Разбивает текст на чанки по предложениям."""
    sents = re.split(r'(?<=[.!?])\s+', text)
    chunks = []
    current = []
    current_len = 0
    chunk_id = 0
    for s in sents:
        s = s.strip()
        if not s:
            continue
        s_len = len(s.split())
        if current_len + s_len > CHUNK_SIZE and current:
            chunks.append(('. '.join(current), chunk_id))
            chunk_id += 1
            keep = max(1, len(current) * CHUNK_OVERLAP // CHUNK_SIZE)
            current = current[-keep:]
            current_len = sum(len(s2.split()) for s2 in current)
        current.append(s)
        current_len += s_len
    if current:
        chunks.append(('. '.join(current), chunk_id))
    return chunks


def build_embedding_prefix(product: str, summary: str) -> str:
    """Строит префикс для эмбеддинга: контекст продукта влияет на вектор."""
    parts = []
    if product:
        parts.append(f"[PRODUCT: {product}]")
    if summary:
        parts.append(f"Summary: {summary}")
    return " ".join(parts) + "\n" if parts else ""


def index_documents(docs_dir: Path, default_product: str = "", default_summary: str = ""):
    """Полная переиндексация всех документов с нуля."""
    import chromadb
    from llama_index.embeddings.ollama import OllamaEmbedding

    chroma_client = chromadb.PersistentClient(str(CHROMA_DIR))
    try:
        chroma_client.delete_collection(COLLECTION_NAME)
    except Exception:
        pass
    collection = chroma_client.create_collection(
        name=COLLECTION_NAME,
        metadata={"hnsw:space": "cosine"}
    )

    embed_model = OllamaEmbedding(
        model_name="nomic-embed-text",
        base_url="http://localhost:11434",
        ollama_additional_kwargs={"mirostat": 0},
    )

    embed_prefix = build_embedding_prefix(default_product, default_summary)

    files = sorted(docs_dir.glob('*'))
    total = len(files)

    all_texts = []
    all_metadatas = []
    all_ids = []

    print(f"Индексация {total} файлов...")
    for i, f in enumerate(files, 1):
        if f.suffix.lower() == '.zip':
            continue
        print(f"  [{i}/{total}] {f.name}")
        text = extract_text(f)
        if not text.strip():
            continue
        chunks = split_text(text, f.name)
        for chunk_text, chunk_id in chunks:
            node_id = f"{f.name}_{chunk_id}"
            all_texts.append(chunk_text)
            all_metadatas.append({
                "source": f.name,
                "chunk": chunk_id,
                "product": default_product,
                "summary": default_summary,
            })
            all_ids.append(node_id)

    # --- векторная индексация ---
    batch_size = 32
    print(f"Генерация эмбеддингов для {len(all_texts)} чанков...")
    for i in range(0, len(all_texts), batch_size):
        batch_texts = all_texts[i:i+batch_size]
        batch_metas = all_metadatas[i:i+batch_size]
        batch_ids = all_ids[i:i+batch_size]
        embed_texts = [embed_prefix + t if embed_prefix else t for t in batch_texts]
        embeds = embed_model.get_text_embedding_batch(embed_texts)
        collection.add(
            embeddings=embeds,
            documents=batch_texts,
            metadatas=batch_metas,
            ids=batch_ids,
        )
        print(f"  векторизовано {min(i+batch_size, len(all_texts))}/{len(all_texts)}")

    print(f"\nВекторная БД: {len(all_texts)} чанков сохранено")

    # --- BM25 индекс ---
    bm25_data = {
        "documents": all_texts,
        "metadatas": all_metadatas,
        "ids": all_ids,
    }
    with open(str(CHROMA_DIR / "bm25_data.pkl"), "wb") as f:
        pickle.dump(bm25_data, f)
    print(f"BM25 данные: {len(all_texts)} документов сохранено")

    # Сброс кэша хешей
    current_hashes = {}
    for f in files:
        if f.suffix.lower() == '.zip':
            continue
        current_hashes[f.name] = file_hash(f)
    with open(HASH_CACHE, "w") as f:
        json.dump(current_hashes, f, indent=2)

    log_ingestion("full_index", files=total, chunks=len(all_texts),
                   product=default_product, summary=default_summary)

    print("✅ Полная индексация завершена!")


def index_incremental(product: str = "", summary: str = ""):
    """Инкрементальная индексация: только новые/изменённые файлы."""
    import chromadb
    from llama_index.embeddings.ollama import OllamaEmbedding

    chroma_client = chromadb.PersistentClient(str(CHROMA_DIR))
    try:
        collection = chroma_client.get_collection(COLLECTION_NAME)
    except Exception:
        print("Коллекция не найдена, создаю новую...")
        collection = chroma_client.create_collection(
            name=COLLECTION_NAME,
            metadata={"hnsw:space": "cosine"}
        )

    embed_model = OllamaEmbedding(
        model_name="nomic-embed-text",
        base_url="http://localhost:11434",
        ollama_additional_kwargs={"mirostat": 0},
    )

    embed_prefix = build_embedding_prefix(product, summary)

    # --- загрузка кэша хешей ---
    hashes = {}
    if HASH_CACHE.exists():
        with open(HASH_CACHE) as f:
            hashes = json.load(f)

    # --- сканирование файлов ---
    files = sorted(DATA_PDFS.glob('*'))
    current_hashes = {}
    to_process = []

    for f in files:
        if f.suffix.lower() == '.zip':
            continue
        h = file_hash(f)
        current_hashes[f.name] = h
        if f.name not in hashes or hashes[f.name] != h:
            to_process.append(f)

    # --- удалённые файлы ---
    removed = [name for name in hashes if name not in current_hashes]

    if not to_process and not removed:
        print("Новых или изменённых файлов нет.")
        return

    # --- удаление старых чанков ---
    for f in to_process:
        if f.name in hashes:  # был, изменился
            print(f"  Удаление старых чанков: {f.name}")
            collection.delete(where={"source": f.name})

    for name in removed:
        print(f"  Удаление чанков удалённого файла: {name}")
        collection.delete(where={"source": name})

    # --- обработка новых/изменённых файлов ---
    new_texts = []
    new_metadatas = []
    new_ids = []

    for f in to_process:
        print(f"  Обработка: {f.name}")
        text = extract_text(f)
        if not text.strip():
            continue
        chunks = split_text(text, f.name)
        for chunk_text, chunk_id in chunks:
            node_id = f"{f.name}_{chunk_id}"
            new_texts.append(chunk_text)
            new_metadatas.append({
                "source": f.name,
                "chunk": chunk_id,
                "product": product,
                "summary": summary,
            })
            new_ids.append(node_id)

    # --- эмбеддинги и добавление в ChromaDB ---
    if new_texts:
        batch_size = 32
        print(f"Генерация эмбеддингов для {len(new_texts)} новых чанков...")
        for i in range(0, len(new_texts), batch_size):
            batch_texts = new_texts[i:i+batch_size]
            batch_metas = new_metadatas[i:i+batch_size]
            batch_ids = new_ids[i:i+batch_size]
            embed_texts = [embed_prefix + t if embed_prefix else t for t in batch_texts]
            embeds = embed_model.get_text_embedding_batch(embed_texts)
            collection.add(
                embeddings=embeds,
                documents=batch_texts,
                metadatas=batch_metas,
                ids=batch_ids,
            )
            print(f"  векторизовано {min(i+batch_size, len(new_texts))}/{len(new_texts)}")

    # --- перестроение BM25 из актуальных данных ChromaDB ---
    print("Перестроение BM25 индекса из актуальных данных...")
    all_data = collection.get()
    bm25_data = {
        "documents": all_data["documents"],
        "metadatas": all_data["metadatas"],
        "ids": all_data["ids"],
    }
    with open(str(CHROMA_DIR / "bm25_data.pkl"), "wb") as f:
        pickle.dump(bm25_data, f)
    print(f"BM25: {len(bm25_data['documents'])} документов сохранено")

    # --- обновление кэша хешей ---
    with open(HASH_CACHE, "w") as f:
        json.dump(current_hashes, f, indent=2)

    log_ingestion("incremental_index", processed=len(to_process),
                   removed=len(removed), new_chunks=len(new_texts),
                   product=product, summary=summary)

    print(f"✅ Инкрементальная индексация завершена: "
          f"{len(to_process)} обработано, {len(removed)} удалено")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Гибридная индексация pdf-rag")
    parser.add_argument('--incremental', action='store_true',
                        help='Инкрементальный режим: только новые/изменённые файлы')
    parser.add_argument('--product', default='',
                        help='Название продукта (для метаданных новых файлов)')
    parser.add_argument('--summary', default='',
                        help='Краткое описание (для метаданных новых файлов)')
    args = parser.parse_args()

    deduplicate_to_flat()

    if args.incremental:
        index_incremental(product=args.product, summary=args.summary)
    else:
        files = [f for f in DATA_PDFS.glob('*') if f.suffix.lower() != '.zip']
        if not files:
            print("Нет файлов в data/pdfs/. Помести файлы в data/pdfs_raw/ и запусти сначала.")
            sys.exit(1)
        index_documents(DATA_PDFS,
                        default_product=args.product,
                        default_summary=args.summary)